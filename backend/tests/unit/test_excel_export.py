import json
import uuid
from datetime import datetime, timezone
from decimal import Decimal
from unittest.mock import MagicMock

from app.export.excel import (
    _build_artifacts_sheet,
    _build_code_sheet,
    _build_configurations_sheet,
    _build_data_splits_sheet,
    _build_datasets_sheet,
    _build_dependencies_sheet,
    _build_differences_sheet,
    _build_environments_sheet,
    _build_experiments_sheet,
    _build_investigations_sheet,
    _build_lineage_sheet,
    _build_metrics_sheet,
    _build_pipelines_sheet,
    _build_randomness_sheet,
    _build_reports_sheet,
    _build_reproducibility_sheet,
    _build_runs_sheet,
    build_workbook,
)
from app.investigation.counterfactual import CounterfactualOutcome
from app.lineage.graph import LineageEdge, LineageEdgeType
from openpyxl import Workbook


def _wb():
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def test_build_workbook_produces_all_seventeen_sheets_when_everything_is_empty():
    db = MagicMock()
    db.query.return_value.all.return_value = []
    db.query.return_value.filter.return_value.all.return_value = []
    db.query.return_value.filter_by.return_value.one_or_none.return_value = None

    wb = build_workbook(db)

    expected = {
        "Experiments",
        "Runs",
        "Code",
        "Datasets",
        "Data Splits",
        "Pipelines",
        "Environments",
        "Dependencies",
        "Configurations",
        "Randomness",
        "Metrics",
        "Artifacts",
        "Differences",
        "Reproducibility",
        "Investigations",
        "Lineage",
        "Reports",
    }
    assert set(wb.sheetnames) == expected
    assert "Sheet" not in wb.sheetnames


def test_experiments_sheet_has_header_and_row_per_experiment():
    experiment = MagicMock()
    experiment.id = uuid.uuid4()
    experiment.project_id = uuid.uuid4()
    experiment.name = "sklearn_tabular"
    experiment.description = "desc"
    experiment.workload_type = "sklearn_tabular"
    experiment.created_at = datetime.now(timezone.utc)

    wb = _wb()
    _build_experiments_sheet(wb, [experiment])
    ws = wb["Experiments"]

    assert ws.max_row == 2
    assert ws.cell(row=2, column=3).value == "sklearn_tabular"


def test_runs_sheet_empty_input_is_header_only():
    wb = _wb()
    _build_runs_sheet(wb, [])
    ws = wb["Runs"]
    assert ws.max_row == 1


def _snapshot():
    s = MagicMock()
    s.id = uuid.uuid4()
    s.run_id = uuid.uuid4()
    s.git_commit_sha = "abc123"
    s.is_dirty = False
    s.tree_fingerprint_hash = "deadbeef"
    s.file_count = 3
    return s


def test_code_sheet_populates_from_snapshots():
    wb = _wb()
    _build_code_sheet(wb, [_snapshot()])
    ws = wb["Code"]
    assert ws.max_row == 2
    assert ws.cell(row=2, column=3).value == "abc123"


def test_code_sheet_empty_is_header_only():
    wb = _wb()
    _build_code_sheet(wb, [])
    assert wb["Code"].max_row == 1


def test_datasets_sheet_uses_related_dataset_name():
    version = MagicMock()
    version.id = uuid.uuid4()
    version.dataset.name = "iris"
    version.content_hash = "hash"
    version.row_count = 150
    version.column_count = 5

    wb = _wb()
    _build_datasets_sheet(wb, [version])
    ws = wb["Datasets"]
    assert ws.cell(row=2, column=2).value == "iris"


def test_environments_dependencies_configurations_randomness_metrics_artifacts_empty_are_header_only():
    wb = _wb()
    _build_environments_sheet(wb, [])
    _build_dependencies_sheet(wb, [])
    _build_configurations_sheet(wb, [])
    _build_randomness_sheet(wb, [])
    _build_metrics_sheet(wb, [])
    _build_artifacts_sheet(wb, [])
    for title in ("Environments", "Dependencies", "Configurations", "Randomness", "Metrics", "Artifacts"):
        assert wb[title].max_row == 1


def test_metrics_sheet_converts_decimal_to_float():
    metric = MagicMock()
    metric.id = uuid.uuid4()
    metric.run_id = uuid.uuid4()
    metric.name = "accuracy"
    metric.value = Decimal("0.942")
    metric.captured_at = datetime.now(timezone.utc)

    wb = _wb()
    _build_metrics_sheet(wb, [metric])
    ws = wb["Metrics"]
    assert ws.cell(row=2, column=4).value == 0.942


def test_data_splits_and_pipelines_sheets_always_have_a_fixed_reason_row():
    wb = _wb()
    _build_data_splits_sheet(wb)
    _build_pipelines_sheet(wb)
    for title in ("Data Splits", "Pipelines"):
        ws = wb[title]
        assert ws.max_row == 2
        assert "NOT_AVAILABLE" in ws.cell(row=2, column=1).value


def _diff():
    d = MagicMock()
    d.id = uuid.uuid4()
    d.comparison_id = uuid.uuid4()
    from app.db.models.enums import Confidence, DifferenceCategory, Severity

    d.category = DifferenceCategory.ENVIRONMENT
    d.field = "dependencies.torch"
    d.old_value = "2.5"
    d.new_value = "2.6"
    d.severity = Severity.HIGH
    d.confidence = Confidence.MEDIUM
    d.is_potential_contributor = True
    return d


def test_differences_sheet_populates_from_difference_rows():
    wb = _wb()
    _build_differences_sheet(wb, [_diff()])
    ws = wb["Differences"]
    assert ws.max_row == 2
    assert ws.cell(row=2, column=3).value == "ENVIRONMENT"


def test_reproducibility_sheet_empty_is_header_only():
    wb = _wb()
    _build_reproducibility_sheet(wb, [])
    assert wb["Reproducibility"].max_row == 1


def test_investigations_sheet_folds_in_investigations_and_counterfactuals():
    from app.db.models.enums import DifferenceCategory

    investigation = MagicMock()
    investigation.id = uuid.uuid4()
    investigation.comparison_id = uuid.uuid4()
    investigation.changed_category = DifferenceCategory.ENVIRONMENT
    investigation.changed_field = "dependencies.torch"
    investigation.evidence_strength = MagicMock(value="HIGH")

    counterfactual = MagicMock()
    counterfactual.id = uuid.uuid4()
    counterfactual.comparison_id = uuid.uuid4()
    counterfactual.restored_category = DifferenceCategory.CONFIGURATION
    counterfactual.restored_field = "model.C"
    counterfactual.evidence_strength = MagicMock(value="MODERATE")
    counterfactual.outcome = CounterfactualOutcome.SUPPORTS_CONTRIBUTION

    wb = _wb()
    _build_investigations_sheet(wb, [investigation], [counterfactual])
    ws = wb["Investigations"]

    assert ws.max_row == 3
    kinds = {ws.cell(row=r, column=3).value for r in (2, 3)}
    assert kinds == {"INVESTIGATION", "COUNTERFACTUAL"}


def test_investigations_sheet_empty_stores_is_header_only():
    wb = _wb()
    _build_investigations_sheet(wb, [], [])
    assert wb["Investigations"].max_row == 1


def test_lineage_sheet_rows_reflect_edges_tagged_by_experiment():
    experiment_id = uuid.uuid4()
    from_run, to_run = uuid.uuid4(), uuid.uuid4()
    edge = LineageEdge(from_run_id=from_run, to_run_id=to_run, edge_type=LineageEdgeType.REPRODUCES)

    wb = _wb()
    _build_lineage_sheet(wb, [(experiment_id, edge)])
    ws = wb["Lineage"]

    assert ws.max_row == 2
    assert ws.cell(row=2, column=1).value == str(experiment_id)
    assert ws.cell(row=2, column=4).value == "REPRODUCES"


def test_lineage_sheet_empty_is_header_only():
    wb = _wb()
    _build_lineage_sheet(wb, [])
    assert wb["Lineage"].max_row == 1


def _run(experiment_id):
    from app.db.models.enums import RunStatus, RunType

    run = MagicMock()
    run.id = uuid.uuid4()
    run.experiment_id = experiment_id
    run.run_type = RunType.ORIGINAL
    run.status = RunStatus.COMPLETED
    run.exit_code = 0
    run.started_at = datetime.now(timezone.utc)
    run.finished_at = datetime.now(timezone.utc)
    return run


def test_reports_sheet_row_reflects_report_and_json_round_trips():
    from app.db.models.enums import ComparisonStatus
    from app.reporting.generator import generate_report

    comparison = MagicMock()
    comparison.id = uuid.uuid4()
    comparison.differences = []
    comparison.code_status = ComparisonStatus.SAME
    comparison.dataset_status = ComparisonStatus.SAME
    comparison.environment_status = ComparisonStatus.SAME
    comparison.configuration_status = ComparisonStatus.SAME
    comparison.randomness_status = ComparisonStatus.SAME
    comparison.metrics_status = ComparisonStatus.NOT_COMPARABLE

    experiment = MagicMock()
    experiment.id = uuid.uuid4()
    experiment.name = "sklearn_tabular"
    experiment.description = None
    experiment.workload_type = "sklearn_tabular"

    original_run = _run(experiment.id)

    report = generate_report(
        comparison=comparison,
        experiment=experiment,
        original_run=original_run,
        reproduction_run=_run(experiment.id),
        reproducibility=None,
    )

    wb = _wb()
    _build_reports_sheet(wb, [(comparison, report)])
    ws = wb["Reports"]

    assert ws.cell(row=2, column=2).value == "sklearn_tabular"
    assert ws.cell(row=2, column=4).value == "NOT_COMPARABLE"
    payload = json.loads(ws.cell(row=2, column=5).value)
    assert isinstance(payload, dict)
    assert payload["metric_comparison"]["status"] == "NOT_COMPARABLE"


def test_reports_sheet_empty_is_header_only():
    wb = _wb()
    _build_reports_sheet(wb, [])
    assert wb["Reports"].max_row == 1
